#!/usr/bin/env python3
"""Import a Mandiri e-Statement .xlsx into data/transactions/YYYY-MM.md.

Usage: python3 scripts/import_mandiri_xlsx.py <statement.xlsx> [--force]
Needs: pip install msoffcrypto-tool openpyxl
Password: env STATEMENT_PASSWORD, otherwise prompted. Never store or commit it.

Categories come from data/rules.md (first match wins, like the app); the rest stay
Uncategorized for the user to confirm. Refuses to overwrite an existing month unless --force.
"""
import datetime
import getpass
import io
import json
import os
import re
import sys
import uuid
from pathlib import Path

import msoffcrypto
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MONTHS = {m: i for i, m in enumerate('Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec'.split(), 1)}


def idr(v):
    """'1.234.567,89' -> 1234567 (dot = thousands, comma = decimals)."""
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    return int(str(v).split(',')[0].replace('.', '').strip() or 0)


def load(path):
    return json.loads(re.match(r'^---json\n([\s\S]*?)\n---', Path(path).read_text()).group(1))


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        sys.exit(__doc__)
    src = Path(args[0])
    password = os.environ.get('STATEMENT_PASSWORD') or getpass.getpass('Statement password: ')
    buf = io.BytesIO()
    office = msoffcrypto.OfficeFile(src.open('rb'))
    office.load_key(password=password)
    office.decrypt(buf)
    ws = openpyxl.load_workbook(buf, data_only=True).active

    header = {}
    for row in ws.iter_rows(max_row=17, values_only=True):
        cells = [c for c in row if c is not None]
        for label in ('Dana Masuk', 'Dana Keluar'):
            if any(str(c).startswith(label + '/') for c in cells):
                header[label] = idr(cells[-1])

    rules = load(ROOT / 'data/rules.md')
    now = datetime.datetime.now(datetime.timezone.utc).isoformat()
    rows = []
    for r in ws.iter_rows(min_row=18, values_only=True):
        if not isinstance(r[0], (int, float)):
            continue
        d, mon, y = str(r[4]).split()[:3]
        income, expense = idr(r[15]), idr(r[18])
        kind = 'income' if income else 'expense'
        desc = ' '.join(str(r[7] or '').split())
        rule = next((x for x in rules if x['pattern'].lower() in desc.lower()
                     and x.get('type') in (None, kind)), None)
        rows.append({
            'id': str(uuid.uuid4()), 'date': f'{y}-{MONTHS[mon]:02d}-{int(d):02d}',
            'amount': income or expense, 'currency': 'IDR', 'type': kind,
            'category': rule['category'] if rule else 'Uncategorized',
            'description': desc, 'source_file': src.name, 'created_at': now,
        })

    got_in = sum(t['amount'] for t in rows if t['type'] == 'income')
    got_out = sum(t['amount'] for t in rows if t['type'] == 'expense')
    if (got_in, got_out) != (header.get('Dana Masuk'), header.get('Dana Keluar')):
        sys.exit(f'Totals mismatch: parsed in {got_in:,} / out {got_out:,} vs statement {header}')

    month = rows[0]['date'][:7]
    out = ROOT / f'data/transactions/{month}.md'
    if out.exists() and '--force' not in sys.argv:
        sys.exit(f'{out} exists — merge by hand or pass --force (overwrites reviewed categories!)')
    out.write_text('---json\n' + json.dumps(rows, indent=2, ensure_ascii=False) + '\n---\n')
    uncat = sum(t['category'] == 'Uncategorized' for t in rows)
    print(f'{out.relative_to(ROOT)}: {len(rows)} rows, in {got_in:,} / out {got_out:,} (matches statement), {uncat} uncategorized')


if __name__ == '__main__':
    main()
